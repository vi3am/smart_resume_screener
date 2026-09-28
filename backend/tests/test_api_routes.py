import tempfile
import uuid

from docx import Document
from fastapi.testclient import TestClient

from main import app


client = TestClient(app)


def test_auth_and_job_resume_flow():
    email = f"smoke_{uuid.uuid4().hex[:8]}@example.com"

    register_response = client.post(
        "/auth/register",
        json={"email": email, "password": "Pass123!", "full_name": "Smoke Tester"},
    )
    assert register_response.status_code == 200, register_response.text
    token = register_response.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    login_response = client.post(
        "/auth/login",
        json={"email": email, "password": "Pass123!"},
        headers=headers,
    )
    assert login_response.status_code == 200, login_response.text

    job_response = client.post(
        "/jobs/",
        json={
            "title": "Python Developer",
            "description": "Build APIs with FastAPI and SQLAlchemy",
            "required_skills": "python, fastapi, sql",
        },
        headers=headers,
    )
    assert job_response.status_code == 200, job_response.text
    job_id = job_response.json()["id"]

    empty_job_response = client.post(
        "/jobs/",
        json={"title": "Empty Job", "description": "No resumes yet", "required_skills": "none"},
        headers=headers,
    )
    assert empty_job_response.status_code == 200, empty_job_response.text
    empty_job_id = empty_job_response.json()["id"]

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        doc = Document()
        doc.add_paragraph("John Smith")
        doc.add_paragraph("john.smith@example.com")
        doc.add_paragraph("+1 555 123 4567")
        doc.add_paragraph("Skills: python, fastapi, sql, docker")
        doc.save(tmp.name)
        resume_path = tmp.name

    with open(resume_path, "rb") as resume_file:
        upload_response = client.post(
            f"/jobs/{job_id}/resumes/",
            files={"file": ("john_smith.docx", resume_file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers=headers,
        )
    assert upload_response.status_code == 200, upload_response.text
    resume_id = upload_response.json()["id"]

    list_resumes_response = client.get(f"/jobs/{job_id}/resumes/", headers=headers)
    assert list_resumes_response.status_code == 200, list_resumes_response.text
    assert any(item["id"] == resume_id for item in list_resumes_response.json())

    reparse_response = client.post(f"/jobs/{job_id}/resumes/{resume_id}/reparse", headers=headers)
    assert reparse_response.status_code == 200, reparse_response.text

    rankings_response = client.get(f"/jobs/{job_id}/resumes/rankings", headers=headers)
    assert rankings_response.status_code == 200, rankings_response.text
    assert isinstance(rankings_response.json(), list)

    delete_with_resume_response = client.delete(f"/jobs/{job_id}", headers=headers)
    assert delete_with_resume_response.status_code == 400
    assert "Cannot delete job" in delete_with_resume_response.json()["detail"]

    delete_empty_job_response = client.delete(f"/jobs/{empty_job_id}", headers=headers)
    assert delete_empty_job_response.status_code == 200, delete_empty_job_response.text
    assert delete_empty_job_response.json()["deleted"] is True
